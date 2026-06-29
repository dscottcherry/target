#!/usr/bin/env python3
"""Render a CLM meeting-prep Markdown outline to a styled .docx.

Usage:  python3 render_docx.py clm-outline-<week>.md [clm-outline-<week>.docx]
Needs:  pip install python-docx

Mirrors the teacher's reference layout: title, color-coded section headers
(Treasures=blue, Field Ministry=green, Living as Christians=purple), item
headers, bold runs, blockquoted scriptures (indented italic), and lists.
Plain Markdown — falls back gracefully on anything it doesn't recognize.
"""
import re
import sys

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    sys.exit("python-docx not installed. Run: pip install python-docx  "
             "(or render with: pandoc <md> -o <docx>)")

BLUE = RGBColor(0x2E, 0x75, 0xB6)    # Treasures
GREEN = RGBColor(0x37, 0x56, 0x23)   # Field Ministry
PURPLE = RGBColor(0x70, 0x30, 0xA0)  # Living as Christians
DARK = RGBColor(0x14, 0x14, 0x13)


def section_color(text):
    t = text.upper()
    if "TREASURES" in t:
        return BLUE
    if "FIELD MINISTRY" in t or "APPLY YOURSELF" in t:
        return GREEN
    if "LIVING AS CHRISTIANS" in t:
        return PURPLE
    return None


def add_runs(par, text):
    """Render inline **bold** segments into a paragraph."""
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if seg == "":
            continue
        run = par.add_run(seg)
        run.bold = (i % 2 == 1)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else re.sub(r"\.md$", "", src) + ".docx"
    lines = open(src, encoding="utf-8").read().splitlines()

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)

    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            if level == 1:                       # document title
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = DARK
            elif level == 2:                     # section header (maybe colored)
                p = doc.add_paragraph(); p.space_before = Pt(14)
                r = p.add_run(text); r.bold = True; r.font.size = Pt(14)
                col = section_color(text)
                if col:
                    r.font.color.rgb = col
            elif level == 3:                     # item header
                p = doc.add_paragraph()
                r = p.add_run(text); r.bold = True; r.font.size = Pt(12.5); r.font.color.rgb = DARK
            else:                                # sub-header
                p = doc.add_paragraph()
                r = p.add_run(text); r.bold = True; r.font.size = Pt(11)
            continue

        if line.lstrip().startswith(">"):        # scripture / quote
            text = line.lstrip()[1:].strip()
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.35)
            r = p.add_run(text); r.italic = True
            continue

        lm = re.match(r"^\s*(?:[-*]|(\d+)\.)\s+(.*)$", line)
        if lm:                                   # list item
            style = "List Number" if lm.group(1) else "List Bullet"
            p = doc.add_paragraph(style=style)
            add_runs(p, lm.group(2))
            continue

        p = doc.add_paragraph()                  # plain paragraph
        add_runs(p, line)

    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
